"""
Tests for Seam 4: Multi-Format Export Sinks
(PostgreSQL, Excel, CSV Zip, Parquet, JSON, and DuckDB Mock API)
"""

import pytest
import io
import zipfile
import json
import openpyxl
import pyarrow.parquet as pq
from app.models.schema import GenerateRequest, EntityDefinition, AttributeDefinition
from app.core.engine import DataEngine
from app.exporters.postgres import PostgresExporter
from app.exporters.excel import ExcelExporter
from app.exporters.csv_zip import CsvZipExporter
from app.exporters.parquet import ParquetExporter
from app.exporters.json_doc import JsonDocExporter
from app.exporters.duckdb_api import DuckDBApiManager


@pytest.fixture
def sample_dataset():
    request = GenerateRequest(
        schema_id="test_sch_123",
        locales=["en_US"],
        entities=[
            EntityDefinition(
                name="users",
                record_count=5,
                attributes=[
                    AttributeDefinition(name="id", type="autoincrement", is_primary_key=True),
                    AttributeDefinition(name="name", type="full_name"),
                    AttributeDefinition(name="email", type="email")
                ]
            ),
            EntityDefinition(
                name="orders",
                record_count=8,
                attributes=[
                    AttributeDefinition(name="id", type="autoincrement", is_primary_key=True),
                    AttributeDefinition(name="user_id", type="foreign_key", fk_target="users.id"),
                    AttributeDefinition(name="total_cents", type="numeric_range", params={"min": 500, "max": 9900})
                ]
            )
        ]
    )
    engine = DataEngine()
    return engine.generate(request), request


def test_postgres_ddl_and_dml_exporter(sample_dataset):
    dataset, request = sample_dataset
    sql = PostgresExporter.generate_ddl_and_dml(dataset, request)
    assert 'CREATE TABLE IF NOT EXISTS "users"' in sql
    assert 'CREATE TABLE IF NOT EXISTS "orders"' in sql
    assert 'REFERENCES "users"("id")' in sql
    assert 'INSERT INTO "users"' in sql
    assert 'INSERT INTO "orders"' in sql
    # users must appear before orders in the SQL script
    assert sql.index('CREATE TABLE IF NOT EXISTS "users"') < sql.index('CREATE TABLE IF NOT EXISTS "orders"')


def test_excel_multi_tab_exporter(sample_dataset):
    dataset, _ = sample_dataset
    excel_bytes = ExcelExporter.export_to_bytes(dataset)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "users" in wb.sheetnames
    assert "orders" in wb.sheetnames
    users_sheet = wb["users"]
    # Header + 5 records = 6 rows
    assert users_sheet.max_row == 6


def test_csv_zip_exporter(sample_dataset):
    dataset, _ = sample_dataset
    zip_bytes = CsvZipExporter.export_to_zip(dataset)
    assert len(zip_bytes) > 0

    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
        names = z.namelist()
        assert "users.csv" in names
        assert "orders.csv" in names
        content = z.read("users.csv").decode("utf-8")
        assert "id,name,email" in content


def test_parquet_exporter(sample_dataset):
    dataset, _ = sample_dataset
    parquet_dict = ParquetExporter.export_tables(dataset)
    assert "users" in parquet_dict
    assert "orders" in parquet_dict

    # Read back through PyArrow
    users_table = pq.read_table(io.BytesIO(parquet_dict["users"]))
    assert users_table.num_rows == 5
    assert "id" in users_table.column_names


def test_json_doc_exporter(sample_dataset):
    dataset, _ = sample_dataset
    json_bytes = JsonDocExporter.export_to_json(dataset)
    parsed = json.loads(json_bytes.decode("utf-8"))
    assert "users" in parsed
    assert "orders" in parsed
    assert len(parsed["users"]) == 5


def test_duckdb_mock_api_manager(sample_dataset):
    dataset, request = sample_dataset
    manager = DuckDBApiManager()
    schema_id = request.schema_id or "test_sch_123"

    manager.persist_dataset(schema_id, dataset)

    # Test catalog listing
    catalog = manager.get_catalog(schema_id)
    assert len(catalog["entities"]) == 2

    # Test pagination
    page1 = manager.query_paginated(schema_id, "orders", limit=3, offset=0)
    assert len(page1["data"]) == 3
    assert page1["pagination"]["total_records"] == 8
    assert page1["pagination"]["has_more"] is True

    page2 = manager.query_paginated(schema_id, "orders", limit=3, offset=6)
    assert len(page2["data"]) == 2
    assert page2["pagination"]["has_more"] is False
